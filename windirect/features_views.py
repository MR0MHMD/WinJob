"""صفحات خصوصی ابزارهای دایرکت و گزارش پاسخ‌نامه‌ها."""
from .subscriptions.access import require_feature, require_connection
from io import BytesIO
from uuid import uuid4
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import ProtectedError
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from .models import (DirectRule, Showcase, ConversationForm, FormSubmission, DirectMessageEvent)
from .features_forms import (DirectRuleForm, ShowcaseForm, ConversationFormEditor, CardFormSet, QuestionFormSet, DirectSimulationForm)
from .direct_services import simulate_direct

TOOLS = {
    'direct': (DirectRule, DirectRuleForm, None, 'دایرکت هوشمند', 'قانون دایرکت', 'کلمهٔ فعال‌کننده را به پاسخ متنی، ویترین یا فرم متصل کنید. قوانین به ترتیب اولویت بررسی می‌شوند.'),
    'showcases': (Showcase, ShowcaseForm, CardFormSet, 'ویترین‌ساز', 'ویترین', 'کارت‌های تصویری محصولات را بسازید؛ سپس در یک قانون دایرکت از ویترین استفاده کنید.'),
    'forms': (ConversationForm, ConversationFormEditor, QuestionFormSet, 'فرم‌ساز', 'فرم', 'سؤال‌ها را مرتب کنید؛ پاسخ‌ها مرحله‌به‌مرحله ثبت می‌شوند و خروجی اکسل دارند.'),
}


def _tool(kind):
    """نوع ابزار را فقط از فهرست ثابت بخش‌های پنل می‌پذیرد."""
    if kind not in TOOLS:
        raise Http404
    model, form, formset, title, singular, intro = TOOLS[kind]
    return model, form, formset, {'section': kind, 'kind': kind, 'heading': title, 'singular': singular, 'intro': intro, 'noindex': True}


@login_required(login_url='accounts:login')
def asset_list(request, kind):
    """فهرست و جست‌وجوی ابزارهای حساب‌های مالک."""
    model, _, _, context = _tool(kind)
    items = model.objects.filter(connection__owner=request.user).select_related('connection')
    query = request.GET.get('q', '').strip()[:150]
    if query:
        items = items.filter(name__icontains=query)
    context.update(page_obj=Paginator(items, 12).get_page(request.GET.get('page')), query=query)
    return render(request, 'windirect/tools/list.html', context)


@login_required(login_url='accounts:login')
def asset_edit(request, kind, pk=None):
    """ابزار و فرزندانش را به صورت یک تراکنش ذخیره می‌کند."""
    model, form_class, formset_class, context = _tool(kind)
    obj = get_object_or_404(model, pk=pk, connection__owner=request.user) if pk else model()
    feature = {'direct': 'smart_direct', 'showcases': 'showcases', 'forms': 'forms'}[kind]
    try:
        require_feature(request.user, feature)
    except ValidationError as error:
        messages.error(request, error.messages[0])
        return redirect('windirect:plans')
    form = form_class(request.POST if request.method == 'POST' else None, owner=request.user, instance=obj)
    children = formset_class(request.POST if request.method == 'POST' else None, instance=obj, prefix='items') if formset_class else None
    if request.method == 'POST':
        valid = form.is_valid()
        children_valid = children.is_valid() if children is not None else True
        if valid:
            try:
                require_connection(form.cleaned_data['connection'], feature)
            except ValidationError as error:
                form.add_error(None, error)
                valid = False
        if valid and children_valid:
            with transaction.atomic():
                form.save()
                if children is not None:
                    children.save()
            messages.success(request, 'تغییرات با موفقیت ذخیره شد.')
            return redirect('windirect:asset_list', kind=kind)
    context.update(form=form, children=children, obj=obj, child_limit=10 if kind == 'showcases' else 20)
    return render(request, 'windirect/tools/editor.html', context)


@login_required(login_url='accounts:login')
@require_POST
def asset_toggle(request, kind, pk):
    """وضعیت ابزار متعلق به کاربر را تغییر می‌دهد."""
    model, _, _, _ = _tool(kind)
    with transaction.atomic():
        obj = get_object_or_404(model.objects.select_for_update(), pk=pk, connection__owner=request.user)
        if not obj.is_active:
            try:
                require_connection(obj.connection, {'direct': 'smart_direct', 'showcases': 'showcases', 'forms': 'forms'}[kind])
            except ValidationError as error:
                messages.error(request, error.messages[0])
                return redirect('windirect:plans')
        obj.is_active = not obj.is_active
        obj.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'وضعیت تغییر کرد. فرم‌های از قبل شروع‌شده تا پایان قابل تکمیل هستند.')
    return redirect('windirect:asset_list', kind=kind)


@login_required(login_url='accounts:login')
def asset_delete(request, kind, pk):
    """حذف ابزار بدون سابقه را پس از نمایش صفحهٔ تأیید انجام می‌دهد."""
    model, _, _, context = _tool(kind)
    obj = get_object_or_404(model, pk=pk, connection__owner=request.user)
    if request.method == 'POST':
        try:
            obj.delete()
        except ProtectedError:
            messages.error(request, 'این مورد سابقه یا قانون وابسته دارد؛ برای حفظ اطلاعات، به جای حذف غیرفعالش کنید.')
        else:
            messages.success(request, 'مورد انتخاب‌شده حذف شد.')
        return redirect('windirect:asset_list', kind=kind)
    context['obj'] = obj
    return render(request, 'windirect/tools/delete.html', context)


@login_required(login_url='accounts:login')
def direct_simulation(request):
    """آزمایش پیوستهٔ پیام‌ها با جلوگیری از ارسال مجدد پس از تازه‌سازی."""
    selected = None
    if request.GET.get('event', '').isdigit():
        selected = get_object_or_404(DirectMessageEvent, pk=request.GET['event'], connection__owner=request.user)
    initial = {'message_id': uuid4()}
    if selected:
        initial.update(connection=selected.connection_id, sender_id=selected.sender_id)
    form = DirectSimulationForm(request.POST if request.method == 'POST' else None, owner=request.user, initial=initial)
    if request.method == 'POST' and form.is_valid():
        data = form.cleaned_data
        try:
            event, _ = simulate_direct(connection=data['connection'], sender_id=data['sender_id'], message_id=str(data['message_id']), text=data['text'])
        except ValidationError as error:
            form.add_error(None, error)
        else:
            return redirect(reverse('windirect:direct_simulation') + f'?event={event.pk}')
    history = []
    if selected:
        history = list(DirectMessageEvent.objects.filter(connection=selected.connection, sender_id=selected.sender_id).select_related('rule')[:30])[::-1]
    return render(request, 'windirect/tools/simulation.html', {'form': form, 'history': history, 'section': 'simulation', 'noindex': True})


@login_required(login_url='accounts:login')
def direct_logs(request):
    """گزارش پیام‌های نمونه و نتیجهٔ مسیریابی آن‌ها."""
    items = DirectMessageEvent.objects.filter(connection__owner=request.user).select_related('connection', 'rule')
    return render(request, 'windirect/tools/logs.html', {'page_obj': Paginator(items, 20).get_page(request.GET.get('page')), 'section': 'direct_logs', 'noindex': True})


def _submissions(request, pk):
    """فیلتر مشترک صفحهٔ پاسخ‌ها و خروجی اکسل را اعمال می‌کند."""
    form = get_object_or_404(ConversationForm, pk=pk, connection__owner=request.user)
    status = request.GET.get('status', '')
    items = form.submissions.filter(connection__owner=request.user)
    if status in FormSubmission.Status.values:
        items = items.filter(status=status)
    else:
        status = ''
    return form, items, status


@login_required(login_url='accounts:login')
def submissions(request, pk):
    """نمایش پاسخ‌نامه‌های کامل، ناقص و لغوشدهٔ فرم."""
    form, items, status = _submissions(request, pk)
    return render(request, 'windirect/tools/submissions.html', {'obj': form, 'page_obj': Paginator(items, 20).get_page(request.GET.get('page')), 'status': status, 'statuses': FormSubmission.Status.choices, 'section': 'forms', 'noindex': True})


@login_required(login_url='accounts:login')
def submission_detail(request, pk):
    """سؤال‌های زمان شروع فرم را کنار پاسخ‌های ثبت‌شده نمایش می‌دهد."""
    obj = get_object_or_404(FormSubmission.objects.select_related('form'), pk=pk, connection__owner=request.user)
    rows = [{'prompt': q['prompt'], 'answer': obj.answers[i] if i < len(obj.answers) else 'بدون پاسخ'} for i, q in enumerate(obj.schema['questions'])]
    return render(request, 'windirect/tools/submission_detail.html', {'obj': obj, 'rows': rows, 'section': 'forms', 'noindex': True})


@login_required(login_url='accounts:login')
def export_submissions(request, pk):
    """خروجی واقعی اکسل با متن امن و حفظ صفر ابتدای شماره‌ها ایجاد می‌کند."""
    form, items, _ = _submissions(request, pk)
    # کران شناسه مانع تغییر مجموعهٔ ردیف‌ها با ثبت پاسخ‌نامهٔ جدید هنگام خروجی می‌شود.
    last = items.order_by('-pk').values_list('pk', flat=True).first() or 0
    items = items.filter(pk__lte=last).order_by('pk')
    columns = {}
    for item in items.only('schema').iterator(chunk_size=500):
        for question in item.schema['questions']:
            key = (question['id'], question['prompt'], question['kind'])
            columns.setdefault(key, f"{question['prompt']} [#{question['id']} / {question['kind']}]")
    book = Workbook(write_only=True)
    sheet = book.create_sheet('پاسخ‌های آزمایشی')
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = 'F2'
    sheet.column_dimensions['B'].width = 24
    sheet.column_dimensions['C'].width = 20
    sheet.column_dimensions['D'].width = 26
    sheet.column_dimensions['E'].width = 26

    def append(values, header=False):
        """ورودی کاربر هرگز به فرمول اکسل تبدیل نمی‌شود."""
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub('', str(value or '')))
            cell.data_type = 's'
            if header:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='FD5631')
            cells.append(cell)
        sheet.append(cells)

    append(['شناسه', 'مخاطب آزمایشی', 'وضعیت', 'تاریخ شروع', 'تاریخ تکمیل', *columns.values()], header=True)
    for item in items.iterator(chunk_size=500):
        answers = {(q['id'], q['prompt'], q['kind']): item.answers[i] for i, q in enumerate(item.schema['questions']) if i < len(item.answers)}
        append([item.pk, item.sender_id, item.get_status_display(), item.created_at, item.completed_at or '', *[answers.get(key, '') for key in columns]])
    output = BytesIO()
    book.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="windirect-form-{form.pk}.xlsx"'
    response['Cache-Control'] = 'private, no-store'
    return response
