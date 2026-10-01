"""آزمون مرزبندی کاربران و جریان ابزارهای جدید وین دایرکت."""
from .test_utils import enable_test_trial
from io import BytesIO
from uuid import uuid4
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from .models import (InstagramConnection, DirectRule, Showcase, ShowcaseCard, ConversationForm, FormQuestion, FormSubmission, DirectMessageEvent)
from .features_forms import DirectRuleForm
from .direct_services import simulate_direct


class FeatureTests(TestCase):
    """رفتار واقعی سرویس، فرم‌ها و خروجی اکسل را بررسی می‌کند."""
    def setUp(self):
        self.user = get_user_model().objects.create_user(phone_number='09120002001', password='test')
        self.other = get_user_model().objects.create_user(phone_number='09120002002', password='test')
        enable_test_trial(self.user)
        self.connection = InstagramConnection.objects.create(owner=self.user, username='mine')
        self.foreign = InstagramConnection.objects.create(owner=self.other, username='foreign')
        self.form = ConversationForm.objects.create(connection=self.connection, name='تماس')
        self.name = FormQuestion.objects.create(form=self.form, prompt='نام شما؟', position=1)
        self.phone = FormQuestion.objects.create(form=self.form, prompt='شمارهٔ همراه؟', kind='phone', position=2)
        self.rule = DirectRule.objects.create(connection=self.connection, name='شروع تماس', keyword='فرم', response_type='form', form=self.form)
        self.client.force_login(self.user)

    def send(self, text, sender='visitor', message_id=None):
        """یک پیام مستقل را برای مخاطب نمونه پردازش می‌کند."""
        return simulate_direct(connection=self.connection, sender_id=sender, message_id=message_id or str(uuid4()), text=text)[0]

    def test_form_validation_completion_and_idempotency(self):
        self.assertIn('نام شما', self.send('فرم').response['text'])
        event = self.send('محمد', message_id='same-message')
        again, created = simulate_direct(connection=self.connection, sender_id='visitor', message_id='same-message', text='محمد')
        self.assertFalse(created)
        self.assertEqual(event.pk, again.pk)
        self.assertIn('۱۱ رقم', self.send('اشتباه').response['text'])
        self.assertEqual(FormSubmission.objects.get().answers, ['محمد'])
        self.send('۰۹۱۲۳۴۵۶۷۸۹')
        submission = FormSubmission.objects.get()
        self.assertEqual(submission.status, 'completed')
        self.assertEqual(submission.answers, ['محمد', '09123456789'])
        self.assertIsNotNone(submission.completed_at)

    def test_snapshot_survives_question_edits_and_deactivation(self):
        self.send('فرم')
        self.name.prompt = 'عنوان جدید'
        self.name.save()
        self.phone.delete()
        self.form.is_active = False
        self.form.save()
        self.assertIn('شماره', self.send('علی').response['text'])
        self.send('09123456789')
        self.assertEqual(FormSubmission.objects.get().status, 'completed')
        self.assertEqual(self.send('فرم', sender='new').response['kind'], 'unmatched')

    def test_cancel_and_sender_isolation(self):
        self.send('فرم', sender='one')
        self.send('فرم', sender='two')
        self.send('لغو', sender='one')
        self.assertEqual(FormSubmission.objects.get(sender_id='one').status, 'cancelled')
        self.assertEqual(FormSubmission.objects.get(sender_id='two').answers, [])
        self.send('فرم', sender='one')
        self.assertEqual(FormSubmission.objects.filter(sender_id='one').count(), 2)

    def test_priority_and_exact_matching(self):
        DirectRule.objects.create(connection=self.connection, name='اول', keyword='قيمت', reply_text='پاسخ اول', match_type='contains', priority=1)
        DirectRule.objects.create(connection=self.connection, name='دوم', keyword='قیمت', reply_text='پاسخ دوم', priority=2)
        self.assertEqual(self.send('قیمت').response['text'], 'پاسخ اول')
        self.assertEqual(self.send('لطفا فرم بده').response['kind'], 'unmatched')

    def test_showcase_snapshot_and_disabled_asset(self):
        showcase = Showcase.objects.create(connection=self.connection, name='محصولات')
        card = ShowcaseCard.objects.create(showcase=showcase, title='محصول', image_url='https://example.com/p.jpg', button_title='خرید', button_url='https://example.com/buy')
        DirectRule.objects.create(connection=self.connection, name='ویترین', keyword='خرید', response_type='showcase', showcase=showcase)
        event = self.send('خرید')
        card.title = 'نام جدید'
        card.save()
        event.refresh_from_db()
        self.assertEqual(event.response['cards'][0]['title'], 'محصول')
        showcase.is_active = False
        showcase.save()
        self.assertEqual(self.send('خرید').response['kind'], 'unmatched')

    def test_connection_and_message_identity_validation(self):
        self.send('سلام', message_id='id')
        with self.assertRaises(ValidationError):
            self.send('متفاوت', message_id='id')
        self.connection.status = 'connected'
        self.connection.save()
        with self.assertRaises(ValidationError):
            self.send('فرم')

    def test_rule_form_rejects_foreign_and_wrong_account_assets(self):
        foreign_form = ConversationForm.objects.create(connection=self.foreign, name='خصوصی')
        data = {'connection': self.connection.pk, 'name': 'قانون', 'keyword': 'شروع', 'match_type': 'exact', 'priority': 10, 'response_type': 'form', 'form': foreign_form.pk, 'is_active': 'on'}
        editor = DirectRuleForm(data, owner=self.user)
        self.assertFalse(editor.is_valid())
        own_second = InstagramConnection.objects.create(owner=self.user, username='second')
        wrong_form = ConversationForm.objects.create(connection=own_second, name='حساب دوم')
        data['form'] = wrong_form.pk
        self.assertFalse(DirectRuleForm(data, owner=self.user).is_valid())
        data['form'] = self.form.pk
        self.assertTrue(DirectRuleForm(data, owner=self.user).is_valid())

    def test_private_routes_and_foreign_ids(self):
        foreign_form = ConversationForm.objects.create(connection=self.foreign, name='خصوصی')
        foreign_submission = FormSubmission.objects.create(form=foreign_form, connection=self.foreign, sender_id='secret', schema={'questions': []})
        for name, args in [('asset_edit', ['forms', foreign_form.pk]), ('asset_delete', ['forms', foreign_form.pk]), ('submissions', [foreign_form.pk]), ('export_submissions', [foreign_form.pk]), ('submission_detail', [foreign_submission.pk])]:
            self.assertEqual(self.client.get(reverse('windirect:' + name, args=args)).status_code, 404)
        self.assertEqual(self.client.post(reverse('windirect:asset_toggle', args=['forms', foreign_form.pk])).status_code, 404)
        self.client.logout()
        for name, args in [('asset_list', ['direct']), ('asset_create', ['forms']), ('direct_simulation', []), ('direct_logs', []), ('export_submissions', [self.form.pk])]:
            self.assertEqual(self.client.get(reverse('windirect:' + name, args=args)).status_code, 302)

    def test_pages_use_site_base_and_render(self):
        self.send('فرم')
        submission = FormSubmission.objects.get()
        routes = [('dashboard', []), ('direct_logs', []), ('direct_simulation', []), ('submissions', [self.form.pk]), ('submission_detail', [submission.pk])]
        for kind in ['direct', 'showcases', 'forms']:
            routes.extend([('asset_list', [kind]), ('asset_create', [kind])])
        for name, args in routes:
            with self.subTest(name=name, args=args):
                response = self.client.get(reverse('windirect:' + name, args=args))
                self.assertEqual(response.status_code, 200)
                self.assertTemplateUsed(response, 'parent/base.html')
                self.assertContains(response, 'finder/css/theme.min.css')
        response = self.client.get(reverse('windirect:direct_simulation') + f'?event={DirectMessageEvent.objects.first().pk}')
        self.assertContains(response, 'نام شما')

    def test_editor_saves_children_and_rejects_wrong_child_ids(self):
        data = {'connection': self.connection.pk, 'name': 'جدید', 'completion_text': 'تمام شد', 'cancel_keyword': 'لغو', 'is_active': 'on', 'items-TOTAL_FORMS': 1, 'items-INITIAL_FORMS': 0, 'items-MIN_NUM_FORMS': 1, 'items-MAX_NUM_FORMS': 20, 'items-0-prompt': 'نام؟', 'items-0-kind': 'text', 'items-0-position': 1}
        response = self.client.post(reverse('windirect:asset_create', args=['forms']), data)
        self.assertEqual(response.status_code, 302)
        created = ConversationForm.objects.get(name='جدید')
        self.assertEqual(created.questions.count(), 1)
        data.update({'items-INITIAL_FORMS': 1, 'items-0-id': self.name.pk, 'items-0-form': created.pk, 'items-0-prompt': 'تغییر نام غیرمجاز'})
        response = self.client.post(reverse('windirect:asset_edit', args=['forms', created.pk]), data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'متعلق به این ابزار نیست')
        self.name.refresh_from_db()
        self.assertEqual(self.name.prompt, 'نام شما؟')

    def test_excel_is_real_safe_filtered_and_keeps_phone(self):
        self.send('فرم')
        self.send('=HYPERLINK("https://example.com")')
        self.send('09123456789')
        self.send('فرم', sender='unfinished')
        response = self.client.get(reverse('windirect:export_submissions', args=[self.form.pk]) + '?status=completed')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        rows = list(sheet.rows)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][5].value, '=HYPERLINK("https://example.com")')
        self.assertEqual(rows[1][5].data_type, 's')
        self.assertEqual(rows[1][6].value, '09123456789')
        self.assertTrue(sheet.sheet_view.rightToLeft)

    def test_delete_is_post_and_history_is_protected(self):
        self.send('فرم')
        url = reverse('windirect:asset_delete', args=['forms', self.form.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertTrue(ConversationForm.objects.filter(pk=self.form.pk).exists())
        self.assertEqual(self.client.post(url).status_code, 302)
        self.assertTrue(ConversationForm.objects.filter(pk=self.form.pk).exists())
        self.assertEqual(self.client.get(reverse('windirect:asset_toggle', args=['forms', self.form.pk])).status_code, 405)

    def test_simulation_post_redirect_get_and_duplicate(self):
        payload = {'connection': self.connection.pk, 'sender_id': 'visitor', 'message_id': str(uuid4()), 'text': 'فرم'}
        url = reverse('windirect:direct_simulation')
        first = self.client.post(url, payload)
        self.assertEqual(first.status_code, 302)
        self.client.post(url, payload)
        self.client.get(first.url)
        self.assertEqual(DirectMessageEvent.objects.count(), 1)
        self.assertEqual(FormSubmission.objects.count(), 1)

    def test_choice_and_number_validation(self):
        self.name.kind = 'number'
        self.name.save()
        self.phone.kind = 'choice'
        self.phone.options = 'بله\nخیر'
        self.phone.save()
        self.send('فرم')
        self.send('سن من بیست است')
        self.assertEqual(FormSubmission.objects.get().answers, [])
        self.send('۲۵')
        self.send('شاید')
        self.assertEqual(FormSubmission.objects.get().answers, ['25'])
        self.send('بله')
        self.assertEqual(FormSubmission.objects.get().status, 'completed')

    def test_formset_rejects_empty_duplicate_order_and_excess(self):
        from .features_forms import QuestionFormSet
        base = {'items-TOTAL_FORMS': '0', 'items-INITIAL_FORMS': '0'}
        self.assertFalse(QuestionFormSet(base, instance=self.form, prefix='items').is_valid())
        data = {'items-TOTAL_FORMS': '2', 'items-INITIAL_FORMS': '0'}
        for index in range(2):
            data.update({f'items-{index}-prompt': 'سؤال', f'items-{index}-kind': 'text', f'items-{index}-position': '1'})
        children = QuestionFormSet(data, instance=self.form, prefix='items')
        self.assertFalse(children.is_valid())
        self.assertIn('ترتیب‌ها', str(children.non_form_errors()))
        data['items-TOTAL_FORMS'] = '5000'
        children = QuestionFormSet(data, instance=self.form, prefix='items')
        self.assertFalse(children.is_valid())
        self.assertEqual(len(children.forms), 30)

    def test_email_validation_and_https_cards(self):
        self.name.kind = 'email'
        self.name.save()
        self.send('فرم')
        self.send('invalid-email')
        self.assertEqual(FormSubmission.objects.get().answers, [])
        self.send('hello@example.com')
        self.assertEqual(FormSubmission.objects.get().answers, ['hello@example.com'])
        showcase = Showcase.objects.create(connection=self.connection, name='ویترین')
        card = ShowcaseCard(showcase=showcase, title='کارت', image_url='http://example.com/image.jpg', button_title='خرید', button_url='javascript:alert(1)')
        with self.assertRaises(ValidationError):
            card.full_clean()

    def test_simulator_rejects_foreign_account_and_csrf(self):
        from django.test import Client
        payload = {'connection': self.foreign.pk, 'sender_id': 'visitor', 'message_id': str(uuid4()), 'text': 'فرم'}
        response = self.client.post(reverse('windirect:direct_simulation'), payload)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(DirectMessageEvent.objects.exists())
        secure_client = Client(enforce_csrf_checks=True)
        secure_client.force_login(self.user)
        response = secure_client.post(reverse('windirect:asset_toggle', args=['forms', self.form.pk]))
        self.assertEqual(response.status_code, 403)

    def test_excel_preserves_old_question_columns_after_edit(self):
        self.send('فرم')
        self.send('نام قدیمی')
        self.send('09123456789')
        self.name.prompt = 'نام جدید؟'
        self.name.save()
        self.send('فرم', sender='second')
        self.send('نام جدید', sender='second')
        self.send('09123456780', sender='second')
        response = self.client.get(reverse('windirect:export_submissions', args=[self.form.pk]))
        rows = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(len(rows), 3)
        self.assertEqual(len(rows[0]), 8)
        self.assertIn('نام قدیمی', rows[1])
        self.assertIn('نام جدید', rows[2])
