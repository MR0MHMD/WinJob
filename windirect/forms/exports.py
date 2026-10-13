from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from .queries import _submissions
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


@login_required(login_url="accounts:login")
def export_submissions(request, pk):
    """خروجی واقعی اکسل با متن امن و حفظ صفر ابتدای شماره‌ها ایجاد می‌کند."""
    form, items, _ = _submissions(request, pk)
    # کران شناسه مانع تغییر مجموعهٔ ردیف‌ها با ثبت پاسخ‌نامهٔ جدید هنگام خروجی می‌شود.
    last = items.order_by("-pk").values_list("pk", flat=True).first() or 0
    items = items.filter(pk__lte=last).order_by("pk")
    columns = {}
    for item in items.only("schema").iterator(chunk_size=500):
        for question in item.schema["questions"]:
            key = (question["id"], question["prompt"], question["kind"])
            columns.setdefault(
                key, f"{question['prompt']} [#{question['id']} / {question['kind']}]"
            )
    book = Workbook(write_only=True)
    sheet = book.create_sheet("پاسخ‌های آزمایشی")
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "F2"
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 26
    sheet.column_dimensions["E"].width = 26

    def append(values, header=False):
        """ورودی کاربر هرگز به فرمول اکسل تبدیل نمی‌شود."""
        cells = []
        for value in values:
            cell = WriteOnlyCell(
                sheet, value=ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
            )
            cell.data_type = "s"
            if header:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="FD5631")
            cells.append(cell)
        sheet.append(cells)

    append(
        [
            "شناسه",
            "مخاطب آزمایشی",
            "وضعیت",
            "تاریخ شروع",
            "تاریخ تکمیل",
            *columns.values(),
        ],
        header=True,
    )
    for item in items.iterator(chunk_size=500):
        answers = {
            (q["id"], q["prompt"], q["kind"]): item.answers[i]
            for i, q in enumerate(item.schema["questions"])
            if i < len(item.answers)
        }
        append(
            [
                item.pk,
                item.sender_id,
                item.get_status_display(),
                item.created_at,
                item.completed_at or "",
                *[answers.get(key, "") for key in columns],
            ]
        )
    output = BytesIO()
    book.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="windirect-form-{form.pk}.xlsx"'
    )
    response["Cache-Control"] = "private, no-store"
    return response
